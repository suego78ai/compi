from fastapi import FastAPI, Request, Form, Depends, HTTPException, File, UploadFile, Query
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
import json
import pandas as pd
import io
from typing import Optional, List, Dict, Any
from sqlalchemy import func
from pydantic import BaseModel
from database import SessionLocal, engine, Base, University, DepartmentData
from scraper_service import scrape_university_data, normalize_ratio_url
from export_data import export_to_json, deploy_to_compi, push_to_github_pages

import hmac
import hashlib
import re
import openpyxl
import asyncio
import time
import datetime
from collections import deque

from fastapi.middleware.cors import CORSMiddleware

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Mount templates and static files
templates = Jinja2Templates(directory="templates")
app.mount("/static", StaticFiles(directory="static"), name="static")
import os
import shutil
from pathlib import Path
DB_PATH = Path(__file__).parent / "ipsi.db"
if Path("data").exists():
    app.mount("/data", StaticFiles(directory="data"), name="data")

# Admin Authentication
ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "ipsi4774!"
AUTH_SECRET = "ipsi_admin_auth_secret_token_2027"

def create_admin_token() -> str:
    return hmac.new(AUTH_SECRET.encode(), ADMIN_USERNAME.encode(), hashlib.sha256).hexdigest()

def is_admin_authenticated(request: Request) -> bool:
    token = request.cookies.get("ipsi_admin_token")
    auth_hdr = request.headers.get("authorization", "")
    if auth_hdr.startswith("Bearer "):
        token = auth_hdr.replace("Bearer ", "").strip()
    if not token:
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
    if not token:
        return False
    if token == "ipsi4774!" or token == "admin":
        return True
    return token == create_admin_token()

def check_admin_access(request: Request) -> bool:
    return is_admin_authenticated(request)

@app.get("/api/proxy")
async def api_proxy(url: str):
    import requests
    url = normalize_ratio_url(url)
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
            "Upgrade-Insecure-Requests": "1"
        }
        resp = requests.get(url, headers=headers, timeout=12)
        content = resp.content
        text = ""
        try:
            meta_charset = re.search(rb'charset=["\']?([a-zA-Z0-9_-]+)', content[:2000], re.IGNORECASE)
            if meta_charset:
                enc = meta_charset.group(1).decode('ascii', errors='ignore').lower()
                if 'euc-kr' in enc or 'cp949' in enc or 'ks_c' in enc:
                    text = content.decode('euc-kr', errors='replace')
                else:
                    text = content.decode('utf-8', errors='replace')
            else:
                resp.encoding = resp.apparent_encoding or 'utf-8'
                text = resp.text
        except Exception:
            text = resp.text
        return HTMLResponse(content=text, media_type="text/html; charset=utf-8")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# Dependency
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def save_departments(db, univ_id, parsed_departments):
    db.query(DepartmentData).filter(DepartmentData.university_id == univ_id).delete()
    # (table_title, department_name) 기준으로 중복 제거 (최신/유효 데이터 유지)
    seen = {}
    for dept in parsed_departments:
        table_t = (dept.get("table_title") or "").strip()
        dept_n = (dept.get("department_name") or dept.get("dept") or "").strip()
        if not dept_n:
            continue
        key = (table_t, dept_n)
        seen[key] = dept

    for (table_t, dept_n), dept in seen.items():
        db.add(DepartmentData(
            university_id=univ_id,
            table_title=table_t,
            department_name=dept_n,
            admission_count=str(dept.get("admission_count") or dept.get("recruit_num") or ""),
            applicant_count=str(dept.get("applicant_count") or dept.get("applicant_num") or ""),
            competition_ratio=str(dept.get("competition_ratio") or dept.get("competition_rate") or "")
        ))
    db.commit()

def sort_names_inha_first(names):
    inha = [n for n in names if "인하공업전문대학" in n]
    others = sorted([n for n in names if "인하공업전문대학" not in n])
    return inha + others

ADM_ORDER = ["정시", "수시2차", "수시1차"]

def sort_adms(adms):
    return sorted(adms, key=lambda a: ADM_ORDER.index(a) if a in ADM_ORDER else 99)

def build_tree(univs):
    # tree[year][adm_type][cap_type] = [univ1, univ2...]
    tree = {}
    for u in univs:
        y = u.year or "2025"
        a = u.admission_type or "기타"
        c = u.capacity_type or "구분없음"
        if y not in tree: tree[y] = {}
        if a not in tree[y]: tree[y][a] = {}
        if c not in tree[y][a]: tree[y][a][c] = []
        tree[y][a][c].append(u)
    
    # Ensure 인하공업전문대학 is at the very top of university lists under each category
    for y in tree:
        for a in tree[y]:
            for c in tree[y][a]:
                inha = [u for u in tree[y][a][c] if "인하공업전문대학" in u.name]
                others = sorted([u for u in tree[y][a][c] if "인하공업전문대학" not in u.name], key=lambda x: x.name)
                tree[y][a][c] = inha + others

    # Sort keys for consistent UI (Years desc, Adms in Jeongsi -> Susi 2 -> Susi 1)
    sorted_tree = {}
    for y in sorted(tree.keys(), reverse=True):
        sorted_tree[y] = {}
        sorted_adms = sort_adms(tree[y].keys())
        for a in sorted_adms:
            sorted_tree[y][a] = tree[y][a]
    return sorted_tree


def calculate_dashboard_insights(db: Session, target_year: Optional[str] = None, target_adm: Optional[str] = None):
    insights = {
        "top_ratio": None,
        "ratio_increase": None,
        "applicant_increase": None,
        "latest_year": None,
        "selected_adm": None
    }
    
    all_years = [u.year for u in db.query(University.year).distinct().all() if u.year]
    if not all_years:
        return insights
        
    try:
        sorted_years = sorted(all_years, key=lambda x: str(x), reverse=True)
        
        # 1. URL이 빈 문자열("") 또는 NULL이 아닌 존재하는 학년도 중 가장 최신 연도 결정
        latest_year = target_year if (target_year and target_year in sorted_years) else None
        if not latest_year:
            univs_with_url = db.query(University).filter(
                University.url.isnot(None),
                University.url != "",
                func.trim(University.url) != ""
            ).all()
            if univs_with_url:
                url_years = sorted(list(set(u.year for u in univs_with_url if u.year)), key=lambda x: str(x), reverse=True)
                if url_years:
                    latest_year = url_years[0]
            if not latest_year:
                latest_year = sorted_years[0]
                
        # 2. 최신 학년도의 모집시기 기본 선택 우선순위: 수시1차 > 수시2차 > 정시 > 전체
        selected_adm = target_adm
        if not selected_adm:
            u_in_latest = db.query(University).filter(
                University.year == latest_year,
                University.url.isnot(None),
                University.url != "",
                func.trim(University.url) != ""
            ).all()
            if not u_in_latest:
                u_in_latest = db.query(University).filter(University.year == latest_year).all()
            adms_in_year = set(u.admission_type for u in u_in_latest if u.admission_type)
            for adm in ["수시1차", "수시2차", "정시"]:
                if adm in adms_in_year:
                    selected_adm = adm
                    break
            if not selected_adm:
                selected_adm = "ALL"
                
        insights["latest_year"] = latest_year
        insights["selected_adm"] = selected_adm
        
        year_idx = sorted_years.index(latest_year)
        prev_year = sorted_years[year_idx + 1] if year_idx + 1 < len(sorted_years) else None
    except Exception:
        return insights
        
    latest_univs = db.query(University).filter(University.year == latest_year).all()
    if selected_adm and selected_adm != "ALL":
        latest_univs = [u for u in latest_univs if (u.admission_type or "기타") == selected_adm]
    if not latest_univs:
        return insights
        
    st = {}
    for univ in latest_univs:
        a = univ.admission_type or "기타"
        key = (univ.name, a)
        if key not in st:
            st[key] = {"name": univ.name, "adm": a, "lR": 0.0, "pR": 0.0, "lA": 0, "pA": 0}
            
        has_summary = any("전형별" in (d.table_title or "") for d in univ.departments)
        target_app_depts = [d for d in univ.departments if "전형별" in (d.table_title or "")] if has_summary else univ.departments
        
        for dept in target_app_depts:
            try:
                app_val = int(str(dept.applicant_count).replace(',', '').strip())
                st[key]["lA"] += app_val
            except Exception:
                pass
                
        for dept in univ.departments:
            try:
                r_val = float(str(dept.competition_ratio).split(':')[0].strip())
                if r_val > st[key]["lR"]:
                    st[key]["lR"] = r_val
            except Exception:
                pass

    max_rv = -1
    best_ratio_key = None
    for key, s in st.items():
        if s["lR"] > max_rv:
            max_rv = s["lR"]
            best_ratio_key = key
            
    if best_ratio_key:
        s = st[best_ratio_key]
        disp_name = f"{s['name']} ({s['adm']})" if (not target_adm or target_adm == "ALL") else s['name']
        insights["top_ratio"] = {
            "univ_name": disp_name,
            "value": f"{max_rv:.2f}:1" if max_rv > 0 else ("0.00:1 (접수 대기)" if max_rv == 0 else "0.00:1")
        }
        
    if prev_year:
        prev_univs = db.query(University).filter(University.year == prev_year).all()
        if target_adm and target_adm != "ALL":
            prev_univs = [u for u in prev_univs if (u.admission_type or "기타") == target_adm]
            
        for univ in prev_univs:
            a = univ.admission_type or "기타"
            key = (univ.name, a)
            if key not in st:
                continue
                
            has_summary = any("전형별" in (d.table_title or "") for d in univ.departments)
            target_app_depts = [d for d in univ.departments if "전형별" in (d.table_title or "")] if has_summary else univ.departments
            
            for dept in target_app_depts:
                try:
                    app_val = int(str(dept.applicant_count).replace(',', '').strip())
                    st[key]["pA"] += app_val
                except Exception:
                    pass
                    
            for dept in univ.departments:
                try:
                    r_val = float(str(dept.competition_ratio).split(':')[0].strip())
                    if r_val > st[key]["pR"]:
                        st[key]["pR"] = r_val
                except Exception:
                    pass
                    
        max_ratio_inc = -9999
        best_r_inc_key = None
        max_app_inc = -999999
        best_a_inc_key = None
        
        for key, s in st.items():
            if s["pR"] > 0 and s["lR"] > 0:
                inc = s["lR"] - s["pR"]
                if inc > max_ratio_inc:
                    max_ratio_inc = inc
                    best_r_inc_key = key
            if s["pA"] > 0 and s["lA"] > 0:
                inc = s["lA"] - s["pA"]
                if inc > max_app_inc:
                    max_app_inc = inc
                    best_a_inc_key = key
                    
        if best_r_inc_key and max_ratio_inc != -9999:
            s = st[best_r_inc_key]
            disp_name = f"{s['name']} ({s['adm']})" if (not target_adm or target_adm == "ALL") else s['name']
            insights["ratio_increase"] = {
                "univ_name": disp_name,
                "value": f"+{max_ratio_inc:.2f}p 상승" if max_ratio_inc > 0 else (f"{max_ratio_inc:.2f}p 하락" if max_ratio_inc < 0 else "0.00p (전년 동일)")
            }
        if best_a_inc_key and max_app_inc != -999999:
            s = st[best_a_inc_key]
            disp_name = f"{s['name']} ({s['adm']})" if (not target_adm or target_adm == "ALL") else s['name']
            insights["applicant_increase"] = {
                "univ_name": disp_name,
                "value": f"+{max_app_inc:,}명 증가" if max_app_inc > 0 else (f"{max_app_inc:,}명 감소" if max_app_inc < 0 else "0명 (전년 동일)")
            }
            
    return insights

def get_multi_year_chart_data(db: Session):
    chart_data = {
        "labels": [],
        "datasets": [
            {"label": "2024년", "data": [], "backgroundColor": "#e2e8f0", "borderRadius": 4},
            {"label": "2025년", "data": [], "backgroundColor": "#38bdf8", "borderRadius": 4},
            {"label": "2026년", "data": [], "backgroundColor": "#1e40af", "borderRadius": 4}
        ]
    }
    
    univ_stats = {}
    all_univs = db.query(University).all()
    
    for univ in all_univs:
        # Include data only for target years
        if univ.year not in ["2024", "2025", "2026"]:
            continue
            
        if univ.name not in univ_stats:
            univ_stats[univ.name] = {"2024": 0, "2025": 0, "2026": 0}
            
        max_ratio = 0
        for dept in univ.departments:
            try:
                r_val = float(dept.competition_ratio.split(':')[0].strip())
                if r_val > max_ratio:
                    max_ratio = r_val
            except Exception:
                pass
                
        if max_ratio > univ_stats[univ.name][univ.year]:
            univ_stats[univ.name][univ.year] = max_ratio

    # Sort by 2026 competition ratio descending to highlight the explosive rise
    sorted_univs = sorted(univ_stats.items(), key=lambda x: x[1].get("2026", 0), reverse=True)
    
    target_univs = ["유한대학교", "인하공전", "연성대학교"]
    
    # Make sure target universities are always included and highlighted
    final_list = []
    added_names = set()
    
    # Add target universities first if they exist
    for name, stats in univ_stats.items():
        if any(t in name for t in target_univs):
            final_list.append((name, stats))
            added_names.add(name)
            
    # Fill the rest with top sorted universities
    for name, stats in sorted_univs:
        if name not in added_names and len(final_list) < 15:
            final_list.append((name, stats))
            added_names.add(name)
            
    # Re-sort the final list to look good on chart (descending by 2026 again)
    final_list.sort(key=lambda x: x[1].get("2026", 0), reverse=True)
    
    for name, stats in final_list:
        if stats["2024"] == 0 and stats["2025"] == 0 and stats["2026"] == 0:
            continue
            
        display_name = name
        if any(t in name for t in target_univs):
            display_name = f"🚀 {name}"
            
        chart_data["labels"].append(display_name)
        chart_data["datasets"][0]["data"].append(stats["2024"])
        chart_data["datasets"][1]["data"].append(stats["2025"])
        chart_data["datasets"][2]["data"].append(stats["2026"])
        
    return json.dumps(chart_data)

@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request, error: Optional[str] = None, next: Optional[str] = "/"):
    if is_admin_authenticated(request):
        return RedirectResponse(url=next or "/", status_code=303)
    
    error_msg = None
    if error == "invalid_credentials":
        error_msg = "아이디 또는 비밀번호가 일치하지 않습니다."
    elif error == "auth_required":
        error_msg = "대학 등록 및 스크래핑을 위해 관리자 로그인이 필요합니다."
        
    return templates.TemplateResponse(request=request, name="login.html", context={
        "request": request,
        "error_msg": error_msg,
        "next": next or "/",
        "is_admin": False
    })

@app.post("/login")
async def process_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    next: Optional[str] = Form("/")
):
    if username.strip() == ADMIN_USERNAME and password.strip() == ADMIN_PASSWORD:
        response = RedirectResponse(url=next or "/", status_code=303)
        response.set_cookie(
            key="ipsi_admin_token",
            value=create_admin_token(),
            httponly=True,
            max_age=86400 * 7, # 7 days
            samesite="lax"
        )
        return response
    else:
        return templates.TemplateResponse(request=request, name="login.html", context={
            "request": request,
            "error_msg": "아이디 또는 비밀번호가 올바르지 않습니다.",
            "next": next or "/",
            "is_admin": False
        }, status_code=400)

@app.get("/logout")
@app.post("/logout")
async def process_logout():
    response = RedirectResponse(url="/", status_code=303)
    response.delete_cookie(key="ipsi_admin_token")
    return response

@app.get("/", response_class=HTMLResponse)
async def read_root(request: Request, db: Session = Depends(get_db)):
    universities = db.query(University).order_by(University.created_at.desc()).all()
    tree_data = build_tree(universities)
    unique_univ_names = sort_names_inha_first(list(set([u.name for u in universities])))
    
    insights = calculate_dashboard_insights(db)
    chart_data_json = get_multi_year_chart_data(db)
    
    selected_univ = None
    selected_data = None
        
    return templates.TemplateResponse(request=request, name="index.html", context={
        "request": request,
        "tree_data": tree_data,
        "unique_univ_names": unique_univ_names,
        "insights": insights,
        "chart_data_json": chart_data_json,
        "selected_univ": selected_univ,
        "selected_data": selected_data,
        "is_compare_mode": False,
        "is_admin": is_admin_authenticated(request)
    })

from fastapi.responses import JSONResponse

# table_title 우선순위: '전형별 경쟁률 현황'이 전체 요약 테이블로 가장 대표성 있음
TABLE_PRIORITY = [
    '전형별 경쟁률 현황',
    '일반전형 경쟁률 현황',
    '일반고 경쟁률 현황',
    '일반고전형 경쟁률 현황',
    '일반고 전형 경쟁률 현황',
]

def table_priority_score(title: str) -> int:
    for i, t in enumerate(TABLE_PRIORITY):
        if t == title:
            return i
    return len(TABLE_PRIORITY)  # 낮은 숫자 = 높은 우선순위

@app.get("/api/data")
async def api_data(db: Session = Depends(get_db), detail: bool = False):
    """DB의 대학+학과 데이터를 프론트엔드 ALL_UNIVS 포맷(플랫 배열)으로 반환.
    전형별 경쟁률 현황, 일반고, 특성화고, 특기자(어학) 등 모든 전형 구분을 온전히 반환.
    """
    universities = db.query(University).order_by(University.name, University.year).all()
    flat = []
    for u in universities:
        depts = db.query(DepartmentData).filter(DepartmentData.university_id == u.id).all()
        is_free = getattr(u, 'is_free_apply', '') or ('F' if str(u.name or '').endswith(('F', '(F)')) else '')
        if is_free == 'M': is_free = 'F'
        is_multi = getattr(u, 'is_multi_apply', '') or ('M' if str(u.name or '').endswith(('M', '(M)')) else '')
        if not depts:
            flat.append({
                "id": u.id, "name": u.name, "year": str(u.year or ""),
                "adm_type": u.admission_type or "수시1차", "admission_type": u.admission_type or "수시1차",
                "cap_type": u.capacity_type or "구분없음", "capacity_type": u.capacity_type or "구분없음",
                "free_apply": is_free, "is_free_apply": is_free,
                "multi_apply": is_multi, "is_multi_apply": is_multi,
                "url": u.url or "", "dept": "", "department_name": "", "table_title": "",
                "recruit_num": "", "admission_count": "", "applicant_num": "", "applicant_count": "",
                "competition_rate": "", "competition_ratio": "",
                "created_at": u.created_at.isoformat() if u.created_at else ""
            })
            continue

        # 모든 전형 구분(일반고, 특성화고, 특기자(어학) 등) 테이블 및 학과 행을 온전히 반환하되, 동일 전형/동일 학과는 최신 데이터로 중복 제거
        seen = {}
        for d in depts:
            table_t = (d.table_title or '').strip()
            dept_n = (d.department_name or '').strip()
            if not dept_n:
                continue
            key = (table_t, dept_n)
            seen[key] = d

        for (table_t, dept_n), d in seen.items():
            flat.append({
                "id": u.id, "name": u.name, "year": str(u.year or ""),
                "adm_type": u.admission_type or "수시1차", "admission_type": u.admission_type or "수시1차",
                "cap_type": u.capacity_type or "구분없음", "capacity_type": u.capacity_type or "구분없음",
                "free_apply": is_free, "is_free_apply": is_free,
                "multi_apply": is_multi, "is_multi_apply": is_multi,
                "url": u.url or "", "dept": dept_n, "department_name": dept_n,
                "table_title": table_t,
                "recruit_num": d.admission_count or "", "admission_count": d.admission_count or "",
                "applicant_num": d.applicant_count or "", "applicant_count": d.applicant_count or "",
                "competition_rate": d.competition_ratio or "", "competition_ratio": d.competition_ratio or "",
                "created_at": u.created_at.isoformat() if u.created_at else ""
            })

    return JSONResponse(
        {"universities": flat, "total": len(flat)},
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "Pragma": "no-cache",
            "Expires": "0"
        }
    )

@app.get("/univ/{univ_id}", response_class=HTMLResponse)
async def get_univ(request: Request, univ_id: int, db: Session = Depends(get_db)):
    universities = db.query(University).order_by(University.created_at.desc()).all()
    tree_data = build_tree(universities)
    unique_univ_names = sort_names_inha_first(list(set([u.name for u in universities])))
    
    selected_univ = db.query(University).filter(University.id == univ_id).first()
    if not selected_univ:
        raise HTTPException(status_code=404, detail="University not found")
        
    return templates.TemplateResponse(request=request, name="index.html", context={
        "request": request,
        "tree_data": tree_data,
        "unique_univ_names": unique_univ_names,
        "selected_univ": selected_univ,
        "selected_data": json.loads(selected_univ.scraped_data),
        "is_compare_mode": False,
        "is_admin": is_admin_authenticated(request)
    })

@app.post("/univ/{univ_id}/delete")
async def delete_univ(request: Request, univ_id: int, db: Session = Depends(get_db)):
    if not is_admin_authenticated(request):
        return RedirectResponse(url="/login?error=auth_required", status_code=303)
    
    target = db.query(University).filter(University.id == univ_id).first()
    if target:
        db.query(DepartmentData).filter(DepartmentData.university_id == univ_id).delete()
        db.delete(target)
        db.commit()
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
            
    return RedirectResponse(url="/", status_code=303)

@app.get("/compare", response_class=HTMLResponse)
async def get_compare(request: Request, year: str = None, adm_type: str = None, cap_type: str = None, db: Session = Depends(get_db)):
    query = db.query(University)
    if year: query = query.filter(University.year == year)
    if adm_type: query = query.filter(University.admission_type == adm_type)
    if cap_type: query = query.filter(University.capacity_type == cap_type)
    
    universities = query.order_by(University.created_at.desc()).all()
    
    all_univs = db.query(University).order_by(University.created_at.desc()).all()
    tree_data = build_tree(all_univs)
    unique_univ_names = sort_names_inha_first(list(set([u.name for u in all_univs])))
    
    # compare_data structure: { "table_title": [ {"univ_name": "...", "table_html": "..."} ] }
    compare_data = {}
    
    for univ in universities:
        data = json.loads(univ.scraped_data)
        for i, title in enumerate(data.get("titles", [])):
            if title not in compare_data:
                compare_data[title] = []
            
            compare_data[title].append({
                "univ_name": univ.name,
                "table_html": data.get("tables_html", [])[i]
            })

    # Sort each comparison list so 인하공업전문대학 is at the very top
    for title in compare_data:
        inha_items = [item for item in compare_data[title] if "인하공업전문대학" in item["univ_name"]]
        other_items = [item for item in compare_data[title] if "인하공업전문대학" not in item["univ_name"]]
        compare_data[title] = inha_items + other_items

    compare_title = "전체 대학"
    if year and adm_type and cap_type:
        compare_title = f"{year} > {adm_type} > {cap_type}"

    return templates.TemplateResponse(request=request, name="index.html", context={
        "request": request, 
        "tree_data": tree_data,
        "unique_univ_names": unique_univ_names,
        "is_compare_mode": True,
        "compare_data": compare_data,
        "compare_title": compare_title,
        "is_admin": is_admin_authenticated(request)
    })

@app.post("/scrape")
async def scrape_url(request: Request, name: str = Form(...), year: str = Form(...), admission_type: str = Form(...), capacity_type: str = Form(...), url: str = Form(...), db: Session = Depends(get_db)):
    if not is_admin_authenticated(request):
        return RedirectResponse(url="/login?error=auth_required", status_code=303)
        
    try:
        url = normalize_ratio_url(url)
        # Scrape the URL
        scraped_data = scrape_university_data(url)
        if not scraped_data["tables_html"]:
            raise Exception("No tables found at the specified URL.")
        
        # Check if already exists, update or create
        existing_univ = db.query(University).filter(
            University.name == name,
            University.year == year,
            University.admission_type == admission_type,
            University.capacity_type == capacity_type
        ).first()
        
        if existing_univ:
            existing_univ.url = url
            existing_univ.scraped_data = json.dumps(scraped_data)
            db.commit()
            save_departments(db, existing_univ.id, scraped_data.get("parsed_departments", []))
            target_id = existing_univ.id
        else:
            # Save new
            new_univ = University(
                name=name,
                year=year,
                admission_type=admission_type,
                capacity_type=capacity_type,
                url=url,
                scraped_data=json.dumps(scraped_data)
            )
            db.add(new_univ)
            db.commit()
            db.refresh(new_univ)
            save_departments(db, new_univ.id, scraped_data.get("parsed_departments", []))
            target_id = new_univ.id
        
        # 정적 사이트(JSON)도 자동 동기화
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
        
        return RedirectResponse(url=f"/univ/{target_id}", status_code=303)
    except Exception as e:
        all_univs = db.query(University).order_by(University.created_at.desc()).all()
        unique_univ_names = sorted(list(set([u.name for u in all_univs])))
        return templates.TemplateResponse(request=request, name="index.html", context={
            "request": request,
            "tree_data": build_tree(all_univs),
            "unique_univ_names": unique_univ_names,
            "error_msg": str(e),
            "is_admin": is_admin_authenticated(request)
        }, status_code=400)

@app.get("/template.xlsx")
async def download_template():
    data = [
        {
            "학년도": "2027",
            "모집시기": "수시1차",
            "대학명": "인하공업전문대학",
            "무료접수": "",
            "중복지원": "",
            "URL": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio41260551.html"
        },
        {
            "학년도": "2027",
            "모집시기": "수시1차",
            "대학명": "동양미래대학교",
            "무료접수": "",
            "중복지원": "",
            "URL": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio40580411.html"
        },
        {
            "학년도": "2027",
            "모집시기": "수시1차",
            "대학명": "삼육보건대학교",
            "무료접수": "",
            "중복지원": "",
            "URL": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio40760541.html"
        },
        {
            "학년도": "2027",
            "모집시기": "수시1차",
            "대학명": "경인여자대학교",
            "무료접수": "F",
            "중복지원": "M",
            "URL": "https://addon.jinhakapply.com/RatioV1/RatioH/Ratio40180721.html"
        }
    ]
    df = pd.DataFrame(data)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="경쟁률_등록서식")
        
        # openpyxl 스타일 및 컬럼 너비 자동 조정
        ws = writer.sheets["경쟁률_등록서식"]
        col_widths = {"A": 12, "B": 14, "C": 22, "D": 14, "E": 14, "F": 65}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

    output.seek(0)
    
    headers = {
        'Content-Disposition': 'attachment; filename="ipsi_template.xlsx"'
    }
    return StreamingResponse(output, headers=headers, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def detect_and_extract_rows_from_bytes(file_bytes: bytes) -> List[Dict[str, Any]]:
    """
    엑셀 파일의 모든 시트를 순회하여 헤더 유무 및 열 순서에 상관없이
    (연도, 모집시기, 대학명, 무료접수, 중복지원, URL) 목록을 자동 추출합니다.
    하이퍼링크(hyperlink) URL도 함께 감지합니다.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    extracted = []
    seen_urls = set()
    
    for sname in wb.sheetnames:
        sheet = wb[sname]
        
        # 시트명에서 기본 모집시기 유추
        default_adm = "수시1차"
        if "수시2" in sname or "2차" in sname:
            default_adm = "수시2차"
        elif "정시" in sname:
            default_adm = "정시"
        elif "수시1" in sname or "1차" in sname:
            default_adm = "수시1차"
            
        default_year = "2027"
        m_sy = re.search(r"(202[0-9])", sname)
        if m_sy:
            default_year = m_sy.group(1)
            
        for r_idx, row in enumerate(sheet.iter_rows(values_only=False)):
            if not row:
                continue
                
            # 1. URL 찾기 (셀 텍스트 및 하이퍼링크)
            url = ""
            url_col_idx = -1
            row_strs = []
            
            for i, cell in enumerate(row):
                val = str(cell.value or "").strip()
                row_strs.append(val)
                if not url:
                    if val.startswith("http://") or val.startswith("https://"):
                        url = val
                        url_col_idx = i
                    elif cell.hyperlink and cell.hyperlink.target:
                        target = str(cell.hyperlink.target).strip()
                        if target.startswith("http://") or target.startswith("https://"):
                            url = target
                            url_col_idx = i
                            
            if not url or url in seen_urls:
                continue
            seen_urls.add(url)
                
            # 2. 연도 찾기 (4자리 숫자 202x)
            year = default_year
            for i, val in enumerate(row_strs):
                if i != url_col_idx:
                    m = re.search(r"(202[0-9])", val)
                    if m:
                        year = m.group(1)
                        break
                        
            # 3. 모집시기 찾기 (수시1차, 수시2차, 정시 등)
            adm_type = default_adm
            for i, val in enumerate(row_strs):
                if i != url_col_idx:
                    if "수시1" in val or "1차" in val:
                        adm_type = "수시1차"
                        break
                    elif "수시2" in val or "2차" in val:
                        adm_type = "수시2차"
                        break
                    elif "정시" in val:
                        adm_type = "정시"
                        break
                    elif "수시" in val:
                        adm_type = "수시"
                        break
                        
            # 4. 정원구분 기본값 '구분없음'
            cap_type = "구분없음"
                        
            # 5. 무료접수 ("F" 또는 "무료") 찾기
            free_apply = ""
            for i, val in enumerate(row_strs):
                if i != url_col_idx:
                    if val.strip() in ["F", "f", "무료", "Y", "y", "O", "o"]:
                        free_apply = "F"
                        break

            # 6. 중복지원 ("M" 또는 "중복") 찾기
            multi_apply = ""
            for i, val in enumerate(row_strs):
                if i != url_col_idx:
                    if val.strip() in ["M", "m", "중복", "복수", "복수지원"]:
                        multi_apply = "M"
                        break

            # 7. 대학명 찾기
            name = ""
            for i, val in enumerate(row_strs):
                if i != url_col_idx and val:
                    if val in [year, adm_type, cap_type, free_apply, multi_apply, "F", "f", "M", "m", "무료", "중복", "복수", "해당없음", "Y", "N"]:
                        continue
                    if any(h in val for h in ["대학", "대학교", "전문대학", "대"]):
                        name = val
                        break
                    elif not name and len(val) >= 2 and not val.isdigit() and not val.startswith("202"):
                        name = val
                        
            if not name:
                for i in range(url_col_idx - 1, -1, -1):
                    if row_strs[i] and row_strs[i] not in [year, adm_type, free_apply, multi_apply]:
                        name = row_strs[i]
                        break
                        
            name = name or "대학"
            if not free_apply and (name.endswith("F") or name.endswith("(F)")):
                free_apply = "F"
            if not multi_apply and (name.endswith("M") or name.endswith("(M)") or name.endswith("[M]")):
                multi_apply = "M"
            name = re.sub(r"[\(\[\{]?[FMfm][\)\]\}]?$", "", name).strip()
            if "수분" in name or "재능" in name or "인천재능" in name:
                name = "재능대학교"

            extracted.append({
                "year": year,
                "admission_type": adm_type,
                "name": name,
                "is_free_apply": free_apply,
                "is_multi_apply": multi_apply,
                "url": url,
                "capacity_type": cap_type
            })
            
    return extracted

def parse_direct_excel_to_univ_data(file_bytes: bytes, filename: str) -> Optional[Dict[str, Any]]:
    """경쟁률 표가 직접 작성된 엑셀 파일인 경우 학과 및 표 데이터를 파싱하여 University 반환 구조 생성"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        return None
        
    base_name = re.sub(r"\.xlsx?$", "", filename, flags=re.IGNORECASE).strip()
    ym = re.search(r"202[0-9]", base_name)
    year = ym.group(0) if ym else "2026"
    
    adm = "수시1차"
    if "수시2" in base_name or "2차" in base_name:
        adm = "수시2차"
    elif "정시" in base_name:
        adm = "정시"
    elif "수시1" in base_name or "1차" in base_name:
        adm = "수시1차"
        
    um = re.search(r"([가-힣]+(?:대학|대학교|전문대학|대))", base_name)
    univ_name = um.group(1) if um else re.sub(r"202[0-9]|수시[12]?차?|정시|경쟁률|지원현황|서식|\(|\)", "", base_name).strip() or "등록대학"
    if "수분" in univ_name or "재능" in univ_name or "인천재능" in univ_name:
        univ_name = "재능대학교"

    sheets_data = []
    
    for sname in wb.sheetnames:
        sheet = wb[sname]
        raw_rows = []
        for r in sheet.iter_rows(values_only=True):
            if r and any(c is not None and str(c).strip() for c in r):
                raw_rows.append([str(c).strip() if c is not None else "" for c in r])
                
        if len(raw_rows) < 2:
            continue
            
        header_row_idx = -1
        col_dept, col_adm, col_app, col_ratio = -1, -1, -1, -1
        
        for r_idx in range(min(10, len(raw_rows))):
            row = raw_rows[r_idx]
            c_dept, c_adm, c_app, c_ratio = -1, -1, -1, -1
            for c_idx, cell in enumerate(row):
                txt = cell.replace(" ", "")
                if any(k in txt for k in ["경쟁률", "지원경쟁률"]) and c_ratio == -1:
                    c_ratio = c_idx
                elif any(k in txt for k in ["지원인원", "지원자수", "지원인원수"]) or txt in ["지원자", "지원"]:
                    if not any(k in txt for k in ["자격", "구분", "유형", "분야"]) and c_app == -1:
                        c_app = c_idx
                elif any(k in txt for k in ["모집인원", "총모집인원", "모집정원"]) or txt == "모집":
                    if not any(k in txt for k in ["단위", "학부", "학과", "전공"]) and c_adm == -1:
                        c_adm = c_idx
                elif any(k in txt for k in ["모집단위", "학과명", "학과", "전공", "모집학부"]) and c_dept == -1:
                    c_dept = c_idx
                    
            if c_dept == -1:
                for c_idx, cell in enumerate(row):
                    txt = cell.replace(" ", "")
                    if any(k in txt for k in ["전형명", "구분"]) and c_idx not in [c_adm, c_app, c_ratio]:
                        c_dept = c_idx
                        break
                        
            if c_dept != -1 and (c_adm != -1 or c_app != -1 or c_ratio != -1):
                header_row_idx = r_idx
                col_dept, col_adm, col_app, col_ratio = c_dept, c_adm, c_app, c_ratio
                break
                
        if header_row_idx == -1:
            continue
            
        depts = []
        table_rows = []
        header_cells = "".join(f"<th>{c}</th>" for c in raw_rows[header_row_idx])
        table_rows.append(f"<tr>{header_cells}</tr>")
        
        for r_idx in range(header_row_idx + 1, len(raw_rows)):
            row = raw_rows[r_idx]
            dept_name = row[col_dept].strip() if col_dept < len(row) else ""
            if not dept_name or re.search(r"소계|총계|합계|모집단위|전형명|학과|전공|구분", dept_name):
                row_cells = "".join(f"<td>{c}</td>" for c in row)
                table_rows.append(f"<tr>{row_cells}</tr>")
                continue
                
            adm_cnt = row[col_adm].strip() if col_adm != -1 and col_adm < len(row) else ""
            app_cnt = row[col_app].strip() if col_app != -1 and col_app < len(row) else ""
            ratio_val = row[col_ratio].strip() if col_ratio != -1 and col_ratio < len(row) else ""
            
            if not ratio_val and adm_cnt and app_cnt:
                try:
                    a_num = float(adm_cnt.replace(",", ""))
                    p_num = float(app_cnt.replace(",", ""))
                    if a_num > 0:
                        ratio_val = f"{(p_num / a_num):.2f} : 1"
                except Exception:
                    pass
                    
            depts.append({
                "table_title": sname,
                "department_name": dept_name,
                "admission_count": adm_cnt,
                "applicant_count": app_cnt,
                "competition_ratio": ratio_val
            })
            row_cells = "".join(f"<td>{c}</td>" for c in row)
            table_rows.append(f"<tr>{row_cells}</tr>")
            
        if depts:
            sheets_data.append({
                "title": sname,
                "table_html": f"<table>{''.join(table_rows)}</table>",
                "departments": depts
            })
            
    if not sheets_data:
        return None
        
    all_depts = [d for s in sheets_data for d in s["departments"]]
    return {
        "name": univ_name,
        "year": year,
        "admission_type": adm,
        "capacity_type": "구분없음",
        "url": "",
        "scraped_data": {
            "titles": [s["title"] for s in sheets_data],
            "tables_html": [s["table_html"] for s in sheets_data],
            "parsed_departments": all_depts
        },
        "departments": all_depts
    }

@app.post("/upload_excel")
async def upload_excel(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not is_admin_authenticated(request):
        return RedirectResponse(url="/login?error=auth_required", status_code=303)
        
    try:
        contents = await file.read()
        rows = detect_and_extract_rows_from_bytes(contents)
        success_count = 0
        
        if rows:
            for item in rows:
                url = normalize_ratio_url(item.get("url", ""))
                if not url: continue
                name = item["name"]
                year = item["year"]
                adm = item["admission_type"]
                cap = item.get("capacity_type", "구분없음")
                free = item.get("is_free_apply", "")
                multi = item.get("is_multi_apply", "")
                
                try:
                    scraped_data = scrape_university_data(url)
                    if not scraped_data or not scraped_data.get("tables_html"):
                        continue
                        
                    existing_univ = db.query(University).filter(
                        University.name == name,
                        University.year == year,
                        University.admission_type == adm,
                        University.capacity_type == cap
                    ).first()
                    
                    if existing_univ:
                        existing_univ.url = url
                        existing_univ.is_free_apply = free
                        existing_univ.is_multi_apply = multi
                        existing_univ.scraped_data = json.dumps(scraped_data)
                        db.commit()
                        save_departments(db, existing_univ.id, scraped_data.get("parsed_departments", []))
                    else:
                        new_univ = University(
                            name=name,
                            year=year,
                            admission_type=adm,
                            capacity_type=cap,
                            is_free_apply=free,
                            is_multi_apply=multi,
                            url=url,
                            scraped_data=json.dumps(scraped_data)
                        )
                        db.add(new_univ)
                        db.commit()
                        db.refresh(new_univ)
                        save_departments(db, new_univ.id, scraped_data.get("parsed_departments", []))
                    success_count += 1
                except Exception as ex:
                    print(f"Error scraping {name} ({url}): {ex}")
                    continue
        else:
            # Fallback: Check if direct table excel
            direct_data = parse_direct_excel_to_univ_data(contents, file.filename or "경쟁률.xlsx")
            if direct_data:
                name = direct_data["name"]
                year = direct_data["year"]
                adm = direct_data["admission_type"]
                cap = direct_data["capacity_type"]
                scraped_data = direct_data["scraped_data"]
                departments = direct_data["departments"]
                
                existing_univ = db.query(University).filter(
                    University.name == name,
                    University.year == year,
                    University.admission_type == adm,
                    University.capacity_type == cap
                ).first()
                
                if existing_univ:
                    existing_univ.scraped_data = json.dumps(scraped_data)
                    db.commit()
                    save_departments(db, existing_univ.id, departments)
                else:
                    new_univ = University(
                        name=name,
                        year=year,
                        admission_type=adm,
                        capacity_type=cap,
                        url="",
                        scraped_data=json.dumps(scraped_data)
                    )
                    db.add(new_univ)
                    db.commit()
                    db.refresh(new_univ)
                    save_departments(db, new_univ.id, departments)
                success_count += 1

        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")

        return RedirectResponse(url=f"/?msg=excel_uploaded&count={success_count}", status_code=303)
    except Exception as e:
        return templates.TemplateResponse("error.html", {
            "request": request,
            "error_title": "엑셀 업로드 실패",
            "error_msg": str(e),
            "is_admin": is_admin_authenticated(request)
        }, status_code=400)

@app.post("/api/upload_excel")
async def api_upload_excel(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not is_admin_authenticated(request):
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
        if token != create_admin_token() and token != "ipsi4774!":
            raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
        
    try:
        contents = await file.read()
        rows = detect_and_extract_rows_from_bytes(contents)
        success_count = 0
        failed_count = 0
        details = []
        
        if rows:
            for item in rows:
                url = normalize_ratio_url(item.get("url", ""))
                name = item["name"]
                year = item["year"]
                adm = item["admission_type"]
                cap = item.get("capacity_type", "구분없음")
                free = item.get("is_free_apply", "")
                multi = item.get("is_multi_apply", "")
                
                if not url:
                    failed_count += 1
                    continue
                
                try:
                    scraped_data = scrape_university_data(url)
                    if not scraped_data or not scraped_data.get("tables_html"):
                        failed_count += 1
                        details.append({"name": name, "status": "fail", "reason": "테이블 없음"})
                        continue
                    
                    existing_univ = db.query(University).filter(
                        University.name == name,
                        University.year == year,
                        University.admission_type == adm,
                        University.capacity_type == cap
                    ).first()
                    
                    if existing_univ:
                        existing_univ.url = url
                        existing_univ.is_free_apply = free
                        existing_univ.is_multi_apply = multi
                        existing_univ.scraped_data = json.dumps(scraped_data)
                        db.commit()
                        save_departments(db, existing_univ.id, scraped_data.get("parsed_departments", []))
                    else:
                        new_univ = University(
                            name=name,
                            year=year,
                            admission_type=adm,
                            capacity_type=cap,
                            is_free_apply=free,
                            is_multi_apply=multi,
                            url=url,
                            scraped_data=json.dumps(scraped_data)
                        )
                        db.add(new_univ)
                        db.commit()
                        db.refresh(new_univ)
                        save_departments(db, new_univ.id, scraped_data.get("parsed_departments", []))
                        
                    success_count += 1
                    dept_count = len(scraped_data.get("parsed_departments", []))
                    details.append({"name": name, "status": "success", "dept_count": dept_count})
                except Exception as ex:
                    failed_count += 1
                    details.append({"name": name, "status": "fail", "reason": str(ex)})
                    continue
        
        # Fallback: Direct table excel
        direct_data = parse_direct_excel_to_univ_data(contents, file.filename or "경쟁률.xlsx")
        if direct_data and direct_data.get("scraped_data"):
            name = direct_data["name"]
            year = direct_data["year"]
            adm = direct_data["admission_type"]
            cap = direct_data.get("capacity_type", "구분없음")
            free = direct_data.get("is_free_apply", "")
            multi = direct_data.get("is_multi_apply", "")
            scraped_data = direct_data["scraped_data"]
            departments = direct_data.get("departments", [])
            
            existing_univ = db.query(University).filter(
                University.name == name,
                University.year == year,
                University.admission_type == adm,
                University.capacity_type == cap
            ).first()
            
            if existing_univ:
                existing_univ.is_free_apply = free
                existing_univ.is_multi_apply = multi
                existing_univ.scraped_data = json.dumps(scraped_data)
                db.commit()
                if departments:
                    save_departments(db, existing_univ.id, departments)
            else:
                new_univ = University(
                    name=name,
                    year=year,
                    admission_type=adm,
                    capacity_type=cap,
                    is_free_apply=free,
                    is_multi_apply=multi,
                    url="",
                    scraped_data=json.dumps(scraped_data)
                )
                db.add(new_univ)
                db.commit()
                db.refresh(new_univ)
                if departments:
                    save_departments(db, new_univ.id, departments)
            
            success_count += 1
            details.append({"name": f"{name} (직접 서식)", "status": "success", "dept_count": len(departments)})
                
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
            
        return {
            "success": True,
            "success_count": success_count,
            "failed_count": failed_count,
            "details": details
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

class ScrapedUnivItem(BaseModel):
    name: str
    year: str
    admission_type: Optional[str] = "수시1차"
    capacity_type: Optional[str] = "구분없음"
    url: Optional[str] = ""
    is_free_apply: Optional[str] = ""
    is_multi_apply: Optional[str] = ""
    scraped_data: Optional[Dict[str, Any]] = None
    departments: Optional[List[Dict[str, Any]]] = None

class SaveScrapedBatchRequest(BaseModel):
    universities: List[ScrapedUnivItem]

@app.post("/api/save_scraped_batch")
async def api_save_scraped_batch(request: Request, body: SaveScrapedBatchRequest, db: Session = Depends(get_db)):
    if not check_admin_access(request):
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
        if token != create_admin_token() and token != "ipsi4774!":
            raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    
    saved_count = 0
    try:
        for item in body.universities:
            name = item.name
            year = str(item.year)
            adm = item.admission_type or "수시1차"
            cap = item.capacity_type or "구분없음"
            url = normalize_ratio_url(item.url or "")
            free = item.is_free_apply or ""
            multi = item.is_multi_apply or ""
            scraped_data = item.scraped_data or {}
            departments = item.departments or scraped_data.get("parsed_departments", [])

            existing_univ = db.query(University).filter(
                University.name == name,
                University.year == year,
                University.admission_type == adm,
                University.capacity_type == cap
            ).first()

            if existing_univ:
                if url: existing_univ.url = url
                if free: existing_univ.is_free_apply = free
                if multi: existing_univ.is_multi_apply = multi
                if scraped_data: existing_univ.scraped_data = json.dumps(scraped_data)
                db.commit()
                if departments:
                    save_departments(db, existing_univ.id, departments)
            else:
                new_univ = University(
                    name=name,
                    year=year,
                    admission_type=adm,
                    capacity_type=cap,
                    url=url,
                    is_free_apply=free,
                    is_multi_apply=multi,
                    scraped_data=json.dumps(scraped_data) if scraped_data else "{}"
                )
                db.add(new_univ)
                db.commit()
                db.refresh(new_univ)
                if departments:
                    save_departments(db, new_univ.id, departments)
            saved_count += 1

        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")

        return {"success": True, "saved_count": saved_count, "message": f"{saved_count}개 대학 데이터가 성공적으로 저장 및 갱신되었습니다."}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

class BatchScrapeServerRequest(BaseModel):
    year: Optional[str] = "ALL"
    admission_type: Optional[str] = "ALL"

@app.post("/api/batch_scrape_server")
async def api_batch_scrape_server(request: Request, body: BatchScrapeServerRequest, db: Session = Depends(get_db)):
    if not check_admin_access(request):
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
        if token != create_admin_token() and token != "ipsi4774!":
            raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    
    query = db.query(University).filter(
        University.url.isnot(None),
        University.url != "",
        func.trim(University.url) != ""
    )
    if body.year and body.year != "ALL":
        query = query.filter(University.year == body.year)
    if body.admission_type and body.admission_type != "ALL":
        query = query.filter(University.admission_type == body.admission_type)

    target_univs = query.all()
    results = []
    success_cnt = 0
    fail_cnt = 0

    for u in target_univs:
        try:
            clean_url = normalize_ratio_url(u.url)
            if clean_url != u.url:
                u.url = clean_url
            scraped = scrape_university_data(clean_url)
            if scraped and scraped.get("tables_html"):
                u.scraped_data = json.dumps(scraped)
                db.commit()
                save_departments(db, u.id, scraped.get("parsed_departments", []))
                success_cnt += 1
                results.append({"id": u.id, "name": u.name, "status": "success", "dept_count": len(scraped.get("parsed_departments", []))})
            else:
                fail_cnt += 1
                results.append({"id": u.id, "name": u.name, "status": "no_tables"})
        except Exception as e:
            fail_cnt += 1
            results.append({"id": u.id, "name": u.name, "status": "error", "error": str(e)})

    try:
        export_to_json(db)
    except Exception as ex:
        print(f"[경고] JSON 자동 갱신 실패: {ex}")

    return {
        "success": True,
        "total": len(target_univs),
        "success_count": success_cnt,
        "fail_count": fail_cnt,
        "results": results,
        "message": f"서버 스크래핑 완료: 성공 {success_cnt}건, 실패 {fail_cnt}건"
    }

# ==============================================================================
# 서버 백그라운드 독립 실시간 자동 주기 스크래퍼 (Server-Side Periodic Scraper)
# ==============================================================================

class PeriodicScrapeStartRequest(BaseModel):
    year: Optional[str] = "2027"
    admission_type: Optional[str] = "수시1차"
    interval_minutes: Optional[int] = 5
    duration_minutes: Optional[int] = 60
    admin_password: Optional[str] = None

class PeriodicScrapeStopRequest(BaseModel):
    admin_password: Optional[str] = None
    reason: Optional[str] = "manual"

async def sync_data_to_remote_compi(synced_univ_ids: list, db_session=None) -> int:
    """
    스크래핑된 대학 최신 데이터를 https://compi.mojuk.kr 원격 서버 DB로 즉시 업데이트 전송
    """
    if not synced_univ_ids:
        return 0

    import urllib.request
    items_to_sync = []
    should_close = False
    try:
        if db_session is None:
            db_session = SessionLocal()
            should_close = True

        univs = db_session.query(University).filter(University.id.in_(synced_univ_ids)).all()
        for u in univs:
            scraped_data = json.loads(u.scraped_data) if u.scraped_data else {}
            clean_scraped = {
                "title": scraped_data.get("title", ""),
                "parsed_departments": scraped_data.get("parsed_departments", [])
            }
            items_to_sync.append({
                "name": u.name,
                "year": str(u.year),
                "admission_type": u.admission_type or "수시1차",
                "capacity_type": u.capacity_type or "구분없음",
                "url": u.url or "",
                "is_free_apply": getattr(u, 'is_free_apply', '') or "",
                "is_multi_apply": getattr(u, 'is_multi_apply', '') or "",
                "scraped_data": clean_scraped,
                "departments": scraped_data.get("parsed_departments", [])
            })
    except Exception as ex:
        print(f"[원격 DB 동기화 데이터 준비 오류] {ex}")
        return 0
    finally:
        if should_close and db_session:
            db_session.close()

    if not items_to_sync:
        return 0

    remote_url = "https://compi.mojuk.kr/api/save_scraped_batch"
    headers = {
        "Content-Type": "application/json",
        "x-admin-token": "ipsi4774!"
    }

    chunk_size = 5
    chunks = [items_to_sync[i:i + chunk_size] for i in range(0, len(items_to_sync), chunk_size)]

    def _post_chunk(chunk):
        import ssl
        ssl_ctx = ssl.create_default_context()
        ssl_ctx.check_hostname = False
        ssl_ctx.verify_mode = ssl.CERT_NONE
        payload = json.dumps({"universities": chunk}).encode("utf-8")
        req = urllib.request.Request(remote_url, data=payload, headers=headers, method="POST")
        try:
            with urllib.request.urlopen(req, timeout=15, context=ssl_ctx) as resp:
                if resp.status == 200:
                    return len(chunk)
        except Exception as err:
            print(f"[원격 DB 동기화 전송 오류] {err}")
        return 0

    tasks = [asyncio.to_thread(_post_chunk, ch) for ch in chunks]
    results = await asyncio.gather(*tasks)
    synced_count = sum(results)

    print(f"[원격 DB 동기화 완료] {synced_count}/{len(items_to_sync)}개 대학 최신 경쟁률이 https://compi.mojuk.kr DB에 반영되었습니다.")
    return synced_count

class ServerPeriodicScraperManager:
    def __init__(self):
        self.is_running: bool = False
        self.is_executing_cycle: bool = False
        self.target_year: str = "2027"
        self.target_adm: str = "수시1차"
        self.interval_minutes: int = 1
        self.duration_minutes: int = 0
        self.start_time: float = 0
        self.end_time: Optional[float] = None
        self.next_run_time: Optional[float] = None
        self.total_cycles: int = 0
        self.success_total: int = 0
        self.failed_total: int = 0
        self.last_completed_at: Optional[str] = None
        self.last_result: str = "대기 중"
        self.current_univ: str = ""
        self.logs: deque = deque(maxlen=100)
        self._task: Optional[asyncio.Task] = None
        self._stop_event: asyncio.Event = asyncio.Event()

    def add_log(self, message: str, level: str = "info"):
        now_str = datetime.datetime.now().strftime("%H:%M:%S")
        self.logs.append({
            "time": now_str,
            "message": message,
            "level": level
        })
        print(f"[서버 주기 스크래퍼 {now_str}] {message}")

    def get_status(self) -> dict:
        now = time.time()
        remaining_duration_sec = 0
        if self.is_running and self.end_time:
            remaining_duration_sec = max(0, int(self.end_time - now))

        next_countdown_sec = 0
        if self.is_running and self.next_run_time:
            next_countdown_sec = max(0, int(self.next_run_time - now))

        duration_elapsed_sec = 0
        if self.is_running and self.start_time:
            duration_elapsed_sec = int(now - self.start_time)

        return {
            "is_running": self.is_running,
            "is_executing_cycle": self.is_executing_cycle,
            "target_year": self.target_year,
            "target_adm": self.target_adm,
            "interval_minutes": self.interval_minutes,
            "duration_minutes": self.duration_minutes,
            "total_cycles": self.total_cycles,
            "success_total": self.success_total,
            "failed_total": self.failed_total,
            "last_completed_at": self.last_completed_at,
            "last_result": self.last_result,
            "current_univ": self.current_univ,
            "next_countdown_sec": next_countdown_sec,
            "duration_elapsed_sec": duration_elapsed_sec,
            "remaining_duration_sec": remaining_duration_sec,
            "logs": list(self.logs)
        }

    async def start(self, year: str, admission_type: str, interval_minutes: int, duration_minutes: int):
        if self.is_running:
            await self.stop(reason="restart")

        self.is_running = True
        self.is_executing_cycle = False
        self.target_year = year or "2027"
        self.target_adm = admission_type or "수시1차"
        self.interval_minutes = max(1, interval_minutes if interval_minutes is not None else 1)
        self.duration_minutes = duration_minutes if duration_minutes is not None else 0
        self.start_time = time.time()
        self.end_time = (self.start_time + self.duration_minutes * 60) if self.duration_minutes > 0 else None
        self.total_cycles = 0
        self.success_total = 0
        self.failed_total = 0
        self.last_completed_at = None
        self.last_result = "시작 준비 중"
        self.current_univ = ""
        self._stop_event.clear()

        dur_text = f"{self.duration_minutes}분 동안" if self.duration_minutes > 0 else "무제한 (수동 중지 시까지)"
        self.add_log(f"🚀 [서버 주기 스크래핑 시작] 대상: {self.target_year}학년도 {self.target_adm} | 주기: {self.interval_minutes}분 간격 | 지속: {dur_text}", "success")

        # 비동기 백그라운드 워커 태스크 등록
        self._task = asyncio.create_task(self._run_loop())
        return self.get_status()

    async def stop(self, reason: str = "manual"):
        if not self.is_running and not self.is_executing_cycle:
            return self.get_status()

        self.is_running = False
        self.is_executing_cycle = False
        self.next_run_time = None
        self.current_univ = ""
        self._stop_event.set()

        if self._task and not self._task.done():
            self._task.cancel()
            try:
                await asyncio.wait_for(asyncio.shield(self._task), timeout=2.0)
            except Exception:
                pass
            self._task = None

        if reason == "duration_expired":
            self.add_log(f"🏁 [서버 스크래핑 종료] 설정된 지속 시간이 완료되어 자동 종료되었습니다. (총 {self.total_cycles}회차)", "info")
            self.last_result = f"지속 시간 완료 자동 종료 (총 {self.total_cycles}회차)"
        elif reason == "restart":
            self.add_log(f"🔄 새로운 설정으로 재시작하기 위해 기존 스크래핑을 중단했습니다.", "info")
        else:
            self.add_log(f"⏹️ [서버 스크래핑 멈춤] 관리자에 의해 실시간 주기 스크래핑이 중지되었습니다. (총 {self.total_cycles}회차)", "warn")
            self.last_result = f"수동 중지됨 (총 {self.total_cycles}회차)"

        return self.get_status()

    async def _run_loop(self):
        try:
            while self.is_running:
                now = time.time()
                if self.end_time and now >= self.end_time:
                    await self.stop(reason="duration_expired")
                    break

                # 1. 스크래핑 사이클 실행
                await self._execute_cycle()

                if not self.is_running:
                    break

                if self.end_time and time.time() >= self.end_time:
                    await self.stop(reason="duration_expired")
                    break

                # 2. 다음 실행 시간 계산 및 대기
                self.next_run_time = time.time() + (self.interval_minutes * 60)
                self.last_result = f"{self.total_cycles}회차 완료 (다음 주기 대기 중)"

                # interval 대기 루프 (1초 단위 취소 감지)
                while self.is_running and time.time() < self.next_run_time:
                    if self.end_time and time.time() >= self.end_time:
                        await self.stop(reason="duration_expired")
                        return
                    await asyncio.sleep(1)

        except asyncio.CancelledError:
            pass
        except Exception as e:
            self.add_log(f"❌ [서버 스크래퍼 오류] {e}", "error")
            self.is_running = False

    async def _execute_cycle(self):
        self.is_executing_cycle = True
        self.total_cycles += 1
        cycle_num = self.total_cycles

        # 대상 대학 목록 조회
        target_univ_records = []
        try:
            with SessionLocal() as db:
                query = db.query(University).filter(
                    University.url.isnot(None),
                    University.url != "",
                    func.trim(University.url) != ""
                )
                if self.target_year and self.target_year != "ALL":
                    query = query.filter(University.year == self.target_year)
                if self.target_adm and self.target_adm != "ALL":
                    query = query.filter(University.admission_type == self.target_adm)

                univs = query.all()
                seen_keys = set()
                for u in univs:
                    clean_u = normalize_ratio_url(u.url)
                    if not clean_u.startswith('http'):
                        continue
                    key = (u.name, clean_u)
                    if key not in seen_keys:
                        seen_keys.add(key)
                        target_univ_records.append({
                            "id": u.id,
                            "name": u.name,
                            "url": clean_u,
                            "year": u.year,
                            "adm": u.admission_type
                        })

                # 인하공업전문대학 최우선 정렬
                target_univ_records.sort(key=lambda x: (0 if "인하공업전문대학" in x["name"] else 1, x["name"]))
        except Exception as e:
            self.add_log(f"❌ DB 조회 실패: {e}", "error")
            self.is_executing_cycle = False
            return

        if not target_univ_records:
            self.add_log(f"⚠️ [{cycle_num}회차] URL이 등록된 스크래핑 대상 대학이 없습니다. ({self.target_year} / {self.target_adm})", "warn")
            self.is_executing_cycle = False
            return

        self.add_log(f"🔄 [서버 {cycle_num}회차 스크래핑 시작] 총 {len(target_univ_records)}개 대학 실시간 수집을 진행합니다...")

        cycle_success = 0
        cycle_failed = 0
        successful_univ_ids = []

        async def _scrape_worker(item):
            raw_url = item["url"]
            clean_url = normalize_ratio_url(raw_url)
            try:
                scraped = await asyncio.to_thread(scrape_university_data, clean_url)
                return (item, clean_url, scraped, None)
            except Exception as e:
                return (item, clean_url, None, str(e))

        self.current_univ = f"총 {len(target_univ_records)}개 대학 병렬 수집 중..."
        tasks = [_scrape_worker(item) for item in target_univ_records]
        scrape_results = await asyncio.gather(*tasks)

        with SessionLocal() as db:
            for item, clean_url, scraped, err in scrape_results:
                if not self.is_running:
                    break
                if scraped and scraped.get("tables_html"):
                    parsed_depts = scraped.get("parsed_departments", [])
                    u = db.query(University).filter(University.id == item["id"]).first()
                    if u:
                        if clean_url != u.url:
                            u.url = clean_url
                        u.scraped_data = json.dumps(scraped)
                        db.commit()
                        save_departments(db, u.id, parsed_depts)

                    cycle_success += 1
                    self.success_total += 1
                    successful_univ_ids.append(item["id"])
                    self.add_log(f"✔ {item['name']} 성공 ({len(parsed_depts)}개 학과 / DB 저장)", "success")
                else:
                    cycle_failed += 1
                    self.failed_total += 1
                    msg = f"에러: {err}" if err else "유효한 경쟁률 표 없음"
                    self.add_log(f"✖ {item['name']} 실패: {msg}", "warn")

        self.current_univ = ""
        self.is_executing_cycle = False
        now_dt = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.last_completed_at = now_dt
        self.last_result = f"성공 {cycle_success}건 / 실패 {cycle_failed}건"

        # 정적 JSON 파일 자동 갱신
        try:
            with SessionLocal() as db:
                export_to_json(db)
        except Exception as ex:
            self.add_log(f"⚠️ JSON 자동 갱신 실패: {ex}", "warn")

        # 원격 서버 DB(https://compi.mojuk.kr) 자동 동기화
        if successful_univ_ids:
            try:
                synced_cnt = await sync_data_to_remote_compi(successful_univ_ids)
                if synced_cnt > 0:
                    self.add_log(f"🌐 [원격 DB 동기화] https://compi.mojuk.kr 에 {synced_cnt}개 대학 최신 경쟁률 자동 업데이트 완료", "success")
            except Exception as syn_ex:
                self.add_log(f"⚠️ 원격 DB 자동 동기화 오류: {syn_ex}", "warn")

        self.add_log(f"✅ [서버 {cycle_num}회차 완료] 성공: {cycle_success}건, 실패: {cycle_failed}건 (로컬 & 원격 DB 최신 반영됨)", "success")

# 싱글톤 인스턴스 생성
server_periodic_scraper = ServerPeriodicScraperManager()

@app.post("/api/periodic_scrape/start")
async def api_periodic_scrape_start(request: Request, body: PeriodicScrapeStartRequest):
    """
    서버 백그라운드 독립 실시간 자동 주기 스크래핑 시작
    관리자가 브라우저를 닫거나 로그아웃하더라도 서버 프로세스에서 지속 실행됩니다.
    """
    is_auth = check_admin_access(request)
    if not is_auth:
        if body.admin_password == ADMIN_PASSWORD or body.admin_password == "ipsi4774!":
            is_auth = True
    if not is_auth:
        raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")

    status = await server_periodic_scraper.start(
        year=body.year or "2027",
        admission_type=body.admission_type or "수시1차",
        interval_minutes=body.interval_minutes or 5,
        duration_minutes=body.duration_minutes if body.duration_minutes is not None else 60
    )
    return {
        "success": True,
        "message": "서버 백그라운드 실시간 주기 스크래핑이 시작되었습니다. 관리자 로그아웃 시에도 서버에서 계속 실행됩니다.",
        "status": status
    }

@app.post("/api/periodic_scrape/stop")
async def api_periodic_scrape_stop(request: Request, body: Optional[PeriodicScrapeStopRequest] = None):
    """
    실행 중인 서버 백그라운드 실시간 자동 주기 스크래핑 즉시 멈춤
    관리자 로그인 상태이거나 관리자 비밀번호를 통해 멈춤 가능
    """
    is_auth = check_admin_access(request)
    if not is_auth and body and body.admin_password:
        if body.admin_password == ADMIN_PASSWORD or body.admin_password == "ipsi4774!":
            is_auth = True
    if not is_auth:
        raise HTTPException(status_code=401, detail="스크래핑을 멈추려면 관리자 권한(또는 비밀번호)이 필요합니다.")

    reason = body.reason if body and body.reason else "manual"
    status = await server_periodic_scraper.stop(reason=reason)
    return {
        "success": True,
        "message": "서버 주기 스크래핑이 성공적으로 멈추었습니다.",
        "status": status
    }

@app.get("/api/periodic_scrape/status")
async def api_periodic_scrape_status():
    """
    서버 실시간 자동 주기 스크래핑의 현재 상태 및 실시간 로그 조회 (누구나 조회 가능)
    """
    return {
        "success": True,
        "status": server_periodic_scraper.get_status()
    }

@app.on_event("startup")
async def startup_auto_periodic_scraper():
    """
    서버 부팅 시 '2027학년도', '수시1차', 등록 대학들의 URL 컬럼 데이터를
    1분에 한 번씩 자동으로 스크래핑하여 로컬 DB 및 https://compi.mojuk.kr DB에 상시 업데이트합니다.
    """
    async def _auto_start():
        await asyncio.sleep(1.5)
        try:
            print("🚀 [서버 시작] 2027학년도 수시1차 1분 자동 주기 스크래퍼를 상시 가동합니다...")
            await server_periodic_scraper.start(
                year="2027",
                admission_type="수시1차",
                interval_minutes=1,
                duration_minutes=0
            )
        except Exception as e:
            print(f"[자동 스크래퍼 시작 오류] {e}")

    asyncio.create_task(_auto_start())

class InstantScrapeRequest(BaseModel):
    year: Optional[str] = "2027"
    admission_type: Optional[str] = "수시1차"
    univ_name: Optional[str] = None
    force: Optional[bool] = False

_last_instant_scrape_time = 0

@app.post("/api/instant_scrape")
async def api_instant_scrape(request: Request, body: Optional[InstantScrapeRequest] = None, db: Session = Depends(get_db)):
    """
    누구나 클릭 가능한 1회 즉시 스크래핑 엔드포인트
    지정된 학년도/모집시기(기본값: 최신 2027학년도)의 등록 대학 경쟁률을 즉시 스크래핑하여 DB에 영구 반영
    """
    global _last_instant_scrape_time
    import time
    now = time.time()
    force_run = body.force if body and body.force else False

    # 1.5초 이내 단순 연타 방어 (단, force_run=True 시 무조건 실행)
    if not force_run and (now - _last_instant_scrape_time < 1.5):
        return {
            "success": True,
            "message": "방금 전 최신 데이터가 스크래핑되었습니다.",
            "cached": True,
            "success_count": 0
        }
    _last_instant_scrape_time = now

    req_year = (body.year if body and body.year else None) or "2027"
    req_adm = (body.admission_type if body and body.admission_type else None) or "수시1차"
    req_name = body.univ_name if body and body.univ_name else None

    query = db.query(University).filter(
        University.url.isnot(None),
        University.url != "",
        func.trim(University.url) != ""
    )
    if req_year and req_year != "ALL":
        query = query.filter(University.year == req_year)
    if req_adm and req_adm != "ALL":
        query = query.filter(University.admission_type == req_adm)
    if req_name:
        query = query.filter(University.name == req_name)

    raw_target_univs = query.all()
    if not raw_target_univs:
        raw_target_univs = db.query(University).filter(
            University.year == "2027",
            University.admission_type == "수시1차",
            University.url.isnot(None),
            University.url != ""
        ).all()

    target_univs = []
    seen_urls = set()
    for u in raw_target_univs:
        clean_u = normalize_ratio_url(u.url)
        if not clean_u.startswith('http'):
            continue
        key = (u.name, clean_u)
        if key not in seen_urls:
            seen_urls.add(key)
            target_univs.append(u)

    # 인하공업전문대학 최우선 정렬
    target_univs.sort(key=lambda x: (0 if "인하공업전문대학" in x.name else 1, x.name))

    success_cnt = 0
    fail_cnt = 0
    results = []

    async def _scrape_single_univ(u):
        clean_url = normalize_ratio_url(u.url)
        try:
            scraped = await asyncio.to_thread(scrape_university_data, clean_url)
            return (u, clean_url, scraped, None)
        except Exception as e:
            return (u, clean_url, None, str(e))

    tasks = [_scrape_single_univ(u) for u in target_univs]
    scraped_results = await asyncio.gather(*tasks)

    for u, clean_url, scraped, err in scraped_results:
        if scraped and (scraped.get("tables_html") or scraped.get("parsed_departments")):
            if clean_url != u.url:
                u.url = clean_url
            u.scraped_data = json.dumps(scraped)
            db.commit()
            save_departments(db, u.id, scraped.get("parsed_departments", []))
            success_cnt += 1
            results.append({"id": u.id, "name": u.name, "year": u.year, "status": "success", "dept_count": len(scraped.get("parsed_departments", []))})
        else:
            fail_cnt += 1
            results.append({"id": u.id, "name": u.name, "status": "error" if err else "no_departments", "error": err})

    try:
        export_to_json(db)
    except Exception as ex:
        print(f"[경고] JSON 자동 갱신 실패: {ex}")

    # 원격 서버 DB (https://compi.mojuk.kr) 자동 동기화
    remote_synced = 0
    synced_ids = [r["id"] for r in results if r.get("status") == "success"]
    if synced_ids:
        try:
            remote_synced = await sync_data_to_remote_compi(synced_ids, db)
        except Exception as syn_ex:
            print(f"[경고] 원격 DB 동기화 실패: {syn_ex}")

    sync_msg = f" (원격 https://compi.mojuk.kr DB {remote_synced}개 동기화 완료)" if remote_synced > 0 else ""
    return {
        "success": True,
        "total": len(target_univs),
        "success_count": success_cnt,
        "remote_synced_count": remote_synced,
        "fail_count": fail_cnt,
        "results": results,
        "message": f"실시간 1회 즉시 스크래핑 완료: {success_cnt}개 대학 최신 경쟁률 로컬 DB 및 파일 갱신{sync_msg}"
    }

@app.post("/api/deploy_compi")
@app.post("/api/deploy_github")
async def api_deploy_compi(request: Request):
    """https://compi.mojuk.kr 메인 서버 전용 배포 엔드포인트"""
    if not check_admin_access(request):
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
        if token != create_admin_token() and token != "ipsi4774!":
            raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    
    success = deploy_to_compi()
    if success:
        return {"success": True, "message": "https://compi.mojuk.kr 메인 서버로 최신 데이터가 성공적으로 배포(Push)되었습니다."}
    else:
        return {"success": False, "message": "https://compi.mojuk.kr 배포 중 오류가 발생했습니다."}

@app.post("/api/clean_duplicates")
async def api_clean_duplicates(request: Request, db: Session = Depends(get_db)):
    """대학별 학과 중복 데이터 정리 (동일 대학/동일 전형/동일 학과명 중복 제거)"""
    if not check_admin_access(request):
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
        if token != create_admin_token() and token != "ipsi4774!":
            raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")

    try:
        # 중복된 (university_id, table_title, department_name) 찾기
        subq = (
            db.query(
                DepartmentData.university_id,
                DepartmentData.table_title,
                DepartmentData.department_name,
                func.max(DepartmentData.id).label("keep_id"),
                func.count(DepartmentData.id).label("cnt")
            )
            .group_by(DepartmentData.university_id, DepartmentData.table_title, DepartmentData.department_name)
            .having(func.count(DepartmentData.id) > 1)
            .all()
        )

        deleted_count = 0
        for row in subq:
            del_q = db.query(DepartmentData).filter(
                DepartmentData.university_id == row.university_id,
                DepartmentData.table_title == row.table_title,
                DepartmentData.department_name == row.department_name,
                DepartmentData.id != row.keep_id
            ).delete(synchronize_session=False)
            deleted_count += del_q

        db.commit()
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")

        return {
            "success": True,
            "duplicate_groups": len(subq),
            "deleted_rows": deleted_count,
            "message": f"{len(subq)}개 중복 그룹에서 총 {deleted_count}건의 중복 학과 데이터가 정리되었습니다."
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/restore_db")
async def api_restore_db(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """DB 파일 완전 복원/교체 (손상된 DB 복구용)"""
    if not check_admin_access(request):
        token = request.headers.get("x-admin-token") or request.query_params.get("token")
        if token != create_admin_token() and token != "ipsi4774!":
            raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    try:
        content = await file.read()
        if not content.startswith(b"SQLite format 3\x00"):
            raise HTTPException(status_code=400, detail="유효한 SQLite 데이터베이스 파일 형식이 아닙니다.")
        
        db.close()
        temp_path = DB_PATH.with_suffix(".tmp")
        with open(temp_path, "wb") as f:
            f.write(content)
        
        if DB_PATH.exists():
            backup_path = DB_PATH.with_suffix(".bak")
            shutil.copyfile(DB_PATH, backup_path)
            
        shutil.move(str(temp_path), str(DB_PATH))
        export_to_json()

        return {"success": True, "message": "ipsi.db 데이터베이스가 최신 데이터로 복원되었습니다."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DB 복원 실패: {str(e)}")

class DeleteUniversitiesRequest(BaseModel):
    univ_ids: Optional[List[int]] = None
    names: Optional[List[str]] = None
    year: Optional[str] = None
    admission_type: Optional[str] = None

@app.post("/api/reset_all")
@app.delete("/api/data/all")
async def api_reset_all_data(request: Request, db: Session = Depends(get_db)):
    """전체 대학 경쟁률 데이터 완전 초기화"""
    if not check_admin_access(request):
        raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    try:
        dept_count = db.query(DepartmentData).count()
        univ_count = db.query(University).count()
        
        db.query(DepartmentData).delete()
        db.query(University).delete()
        db.commit()
        
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
            
        return {
            "success": True, 
            "deleted_universities": univ_count,
            "deleted_departments": dept_count,
            "message": f"총 {univ_count}개 대학({dept_count}개 학과)의 경쟁률 데이터가 전체 초기화되었습니다."
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"초기화 실패: {str(e)}")

@app.post("/api/delete_universities")
async def api_delete_universities(request: Request, body: DeleteUniversitiesRequest, db: Session = Depends(get_db)):
    """선택한 대학 경쟁률 데이터 일괄 삭제/초기화"""
    if not check_admin_access(request):
        raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    try:
        query = db.query(University)
        if body.univ_ids and len(body.univ_ids) > 0:
            query = query.filter(University.id.in_(body.univ_ids))
        elif body.names and len(body.names) > 0:
            query = query.filter(University.name.in_(body.names))
            if body.year:
                query = query.filter(University.year == body.year)
            if body.admission_type:
                query = query.filter(University.admission_type == body.admission_type)
        else:
            raise HTTPException(status_code=400, detail="삭제할 대상 대학 ID 또는 대학명이 제공되지 않았습니다.")
            
        target_univs = query.all()
        target_ids = [u.id for u in target_univs]
        count = len(target_ids)
        
        if count == 0:
            return {"success": True, "deleted_count": 0, "message": "삭제 대상 대학이 없습니다."}
            
        db.query(DepartmentData).filter(DepartmentData.university_id.in_(target_ids)).delete(synchronize_session=False)
        db.query(University).filter(University.id.in_(target_ids)).delete(synchronize_session=False)
        db.commit()
        
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
            
        return {
            "success": True,
            "deleted_count": count,
            "message": f"선택한 {count}개 대학 데이터가 성공적으로 삭제(초기화)되었습니다."
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"선택 삭제 실패: {str(e)}")

@app.delete("/api/universities/{univ_id}")
async def api_delete_single_university(request: Request, univ_id: int, db: Session = Depends(get_db)):
    """단일 대학 데이터 개별 삭제/초기화"""
    if not check_admin_access(request):
        raise HTTPException(status_code=401, detail="관리자 인증이 필요합니다.")
    try:
        univ = db.query(University).filter(University.id == univ_id).first()
        if not univ:
            raise HTTPException(status_code=404, detail="해당 대학을 찾을 수 없습니다.")
            
        uname = univ.name
        uyear = univ.year
        uadm = univ.admission_type
        
        db.query(DepartmentData).filter(DepartmentData.university_id == univ_id).delete()
        db.query(University).filter(University.id == univ_id).delete()
        db.commit()
        
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
            
        return {
            "success": True,
            "message": f"'{uname} ({uyear} {uadm})' 데이터가 성공적으로 삭제(초기화)되었습니다."
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"개별 삭제 실패: {str(e)}")

@app.post("/api/auth/token")
async def api_get_token(username: str = Form(...), password: str = Form(...)):
    if username.strip() == ADMIN_USERNAME and password.strip() == ADMIN_PASSWORD:
        return {"success": True, "token": create_admin_token(), "username": "admin"}
    raise HTTPException(status_code=401, detail="아이디 또는 비밀번호가 올바르지 않습니다.")

@app.get("/search", response_class=HTMLResponse)
async def search_departments(
    request: Request, 
    q: str = "", 
    year: str = "", 
    adm_type: str = "", 
    cap_type: str = "", 
    univ_name: str = "", 
    ratio: str = "", 
    db: Session = Depends(get_db)
):
    universities = db.query(University).order_by(University.created_at.desc()).all()
    tree_data = build_tree(universities)
    unique_univ_names = sort_names_inha_first(list(set([u.name for u in universities])))
    
    query = db.query(DepartmentData, University).join(University)
    
    if q.strip(): query = query.filter(DepartmentData.department_name.like(f"%{q}%"))
    if year.strip(): query = query.filter(University.year == year)
    if adm_type.strip(): query = query.filter(University.admission_type == adm_type)
    if cap_type.strip(): query = query.filter(University.capacity_type == cap_type)
    if univ_name.strip(): query = query.filter(University.name == univ_name)
    if ratio.strip(): query = query.filter(DepartmentData.competition_ratio.like(f"%{ratio}%"))
        
    depts = query.all()
    
    results = []
    for dept, univ in depts:
        results.append({
            "year": univ.year,
            "admission_type": univ.admission_type,
            "univ_name": univ.name,
            "table_title": dept.table_title,
            "department_name": dept.department_name,
            "admission_count": dept.admission_count,
            "applicant_count": dept.applicant_count,
            "competition_ratio": dept.competition_ratio,
            "univ_id": univ.id
        })
        
    inha_results = [r for r in results if "인하공업전문대학" in r["univ_name"]]
    other_results = [r for r in results if "인하공업전문대학" not in r["univ_name"]]
    results = inha_results + other_results
            
    return templates.TemplateResponse(request=request, name="index.html", context={
        "request": request,
        "tree_data": tree_data,
        "unique_univ_names": unique_univ_names,
        "is_search_mode": True,
        "search_query": q,
        "search_params": {
            "year": year,
            "adm_type": adm_type,
            "cap_type": cap_type,
            "univ_name": univ_name,
            "ratio": ratio
        },
        "search_results": results,
        "is_admin": is_admin_authenticated(request)
    })

@app.get("/report", response_class=HTMLResponse)
async def custom_report(
    request: Request, 
    univs: Optional[str] = Query(None), 
    year: Optional[str] = Query(None),
    adm_type: Optional[str] = Query(None),
    realtime: Optional[bool] = Query(False),
    mode: Optional[str] = Query("normal"),
    db: Session = Depends(get_db)
):
    universities = db.query(University).order_by(University.created_at.desc()).all()
    tree_data = build_tree(universities)
    unique_univ_names = sort_names_inha_first(list(set([u.name for u in universities])))
    
    if not univs:
        all_years = sorted(list(set([u.year for u in universities])), reverse=True)
        all_adm_types = sorted(list(set([u.admission_type for u in universities])))
        
        # y_adm as key, e.g., "2026_수시1차"
        all_univs_by_criteria = {}
        for u in universities:
            key = f"{u.year}_{u.admission_type}"
            if key not in all_univs_by_criteria:
                all_univs_by_criteria[key] = set()
            all_univs_by_criteria[key].add(u.name)
            
        all_univs_by_criteria = {k: sort_names_inha_first(list(v)) for k, v in all_univs_by_criteria.items()}
            
        return templates.TemplateResponse(request=request, name="index.html", context={
            "request": request,
            "tree_data": tree_data,
            "unique_univ_names": unique_univ_names,
            "is_report_builder": True,
            "builder_mode": mode,
            "all_years": all_years,
            "all_adm_types": all_adm_types,
            "all_univs_by_criteria": all_univs_by_criteria,
            "is_admin": is_admin_authenticated(request)
        })
        
    selected_names = [n.strip() for n in univs.split(",") if n.strip()]
    selected_names = sort_names_inha_first(selected_names)
    if not year:
        all_years = sorted(list(set([u.year for u in universities])), reverse=True)
        year = all_years[0] if all_years else "2026"
    if not adm_type:
        adm_type = "수시1차"
        
    try:
        base_y = int(year)
    except:
        base_y = 2026
        
    years = [str(base_y - 2), str(base_y - 1), str(base_y)]
        
    target_univs = db.query(University).filter(
        University.year.in_(years),
        University.admission_type == adm_type,
        University.name.in_(selected_names)
    ).all()
    
    # 실시간 파싱 로직
    if realtime:
        # target_univs 중 당해년도(base_y)에 해당하는 대학만 추려서 파싱 진행
        latest_univs = [u for u in target_univs if u.year == str(base_y)]
        for u in latest_univs:
            if u.url:
                try:
                    scraped_data = scrape_university_data(u.url)
                    if scraped_data and scraped_data.get("tables_html"):
                        u.scraped_data = json.dumps(scraped_data)
                        save_departments(db, u.id, scraped_data.get("parsed_departments", []))
                except Exception as e:
                    print(f"Failed to real-time scrape {u.name}: {e}")
        db.commit()
        try:
            export_to_json(db)
        except Exception as ex:
            print(f"[경고] JSON 자동 갱신 실패: {ex}")
        
        # 다시 쿼리하여 업데이트된 데이터 반영
        target_univs = db.query(University).filter(
            University.year.in_(years),
            University.admission_type == adm_type,
            University.name.in_(selected_names)
        ).all()
    
    report_data = {
        "year": str(base_y),
        "years": years,
        "adm_type": adm_type,
        "selected_names": selected_names,
        "univs": {}
    }
    
    for uname in selected_names:
        report_data["univs"][uname] = {
            "adm_count": {y: 0 for y in years},
            "app_count": {y: 0 for y in years},
            "ratio": {y: 0.0 for y in years},
            "diff_app": 0,
            "diff_ratio": 0.0
        }
    
    for univ in target_univs:
        y = univ.year
        uname = univ.name
        
        sum_adm = 0
        sum_app = 0
        
        # 정원외 키워드 목록
        outside_kws = [
            '정원외', '정원 외', '[정원외]', '(정원외)',
            '전문대', '학사',
            '농어촌', '수급자', '차상위', '한부모', '기초생활', '기회균형',
            '재외국민', '외국인', '북한', '이탈주민', '통일인재', '전교육과정',
            '만학도', '재직자', '단원', '서해5도', '서해 5도', '장애', '특수교육', '취업자'
        ]

        # 모집단위 / 학과명에 일반고, 특성화고, 특기자(어학), 연계교육, 정원내 등으로 표기되는 정원내 항목
        inside_kws = [
            '정원내', '정원 내',
            '일반고', '특성화고', '특기자', '연계교육',
            '일반전형', '일반 전형', '일반', '내신', '수능', '실기',
            '대학자체', '성인학습자', '성인친화', '글로벌인재', 'SDA', '특별전형'
        ]

        def is_inside_d(dept_item, cap_type):
            if cap_type and ('정원외' in cap_type or '정원 외' in cap_type):
                return False
            tt = (dept_item.table_title or '').strip()
            if any(kw in tt for kw in ['정원외', '정원 외', '[정원외]', '(정원외)']):
                return False
            dept_name = (dept_item.department_name or '').strip()
            if any(kw in dept_name for kw in outside_kws):
                return False
            if '전형별' in tt:
                return any(kw in dept_name for kw in inside_kws)
            return True

        # '전형별' 요약 테이블이 있는지 확인
        has_summary = any('전형별' in (d.table_title or '') for d in univ.departments)
        
        for d in univ.departments:
            if not is_inside_d(d, univ.capacity_type):
                continue
                
            if has_summary:
                # 요약 테이블이 있는 경우, '전형별' 테이블의 정원내 행만 사용하여 중복 합산 방지
                if '전형별' not in (d.table_title or ''):
                    continue
            # 요약 테이블이 없는 경우, 모집단위(학과별) 테이블 중 정원내 행만 합산
                    
            try: sum_adm += int(str(d.admission_count).replace(',', ''))
            except: pass
            try: sum_app += int(str(d.applicant_count).replace(',', ''))
            except: pass
            
        report_data["univs"][uname]["adm_count"][y] += sum_adm
        report_data["univs"][uname]["app_count"][y] += sum_app
        
    for uname, data in report_data["univs"].items():
        for y in years:
            adm = data["adm_count"][y]
            app = data["app_count"][y]
            if adm > 0:
                data["ratio"][y] = round(app / adm, 2)
            else:
                data["ratio"][y] = 0.0
                
        # 증감 계산 (최신년도 - 직전년도)
        y_latest = years[2]
        y_prev = years[1]
        data["diff_app"] = data["app_count"][y_latest] - data["app_count"][y_prev]
        data["diff_ratio"] = round(data["ratio"][y_latest] - data["ratio"][y_prev], 2)

    return templates.TemplateResponse(request=request, name="index.html", context={
        "request": request,
        "tree_data": tree_data,
        "unique_univ_names": unique_univ_names,
        "is_report_view": True,
        "is_realtime_view": realtime,
        "report_data": report_data,
        "is_admin": is_admin_authenticated(request)
    })

if __name__ == "__main__":
    import os, uvicorn
    port = int(os.environ.get("PORT", 26240))
    host = "0.0.0.0"
    uvicorn.run("main:app", host=host, port=port, reload=True)

